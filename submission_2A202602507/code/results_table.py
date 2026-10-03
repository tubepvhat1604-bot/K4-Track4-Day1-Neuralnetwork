"""results_table.py — lưu kết quả JSON, đọc lại, và điền experiments.xlsx từ mẫu.

save_result / load_results dùng từ Part 2; to_row / write_xlsx dùng ở Part 4.
"""
from __future__ import annotations

import json
import math
import os


def _jsonable(o):
    """Chuyển numpy/torch/tuple/NaN về kiểu JSON ghi được."""
    try:
        import numpy as np
        if isinstance(o, np.generic):
            o = o.item()
    except ImportError:
        pass
    if isinstance(o, float) and (math.isnan(o) or math.isinf(o)):
        return None
    if isinstance(o, dict):
        return {str(k): _jsonable(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_jsonable(v) for v in o]
    return o


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi results/<exp_id>.json gồm cfg, history, summary (KHÔNG ghi best_state: không nộp trọng số)."""
    os.makedirs(results_dir, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    path = os.path.join(results_dir, f"{exp_id}.json")
    payload = {k: result[k] for k in ("cfg", "history", "summary") if k in result}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(_jsonable(payload), f, ensure_ascii=False, indent=2)
    return path


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi results/*.json, sắp theo thứ tự tên file."""
    out = []
    if not os.path.isdir(results_dir):
        return out
    for name in sorted(os.listdir(results_dir)):
        if name.endswith(".json"):
            with open(os.path.join(results_dir, name), encoding="utf-8") as f:
                obj = json.load(f)
            if isinstance(obj, dict) and "cfg" in obj:      # bỏ qua file JSON không phải kết quả chạy
                out.append(obj)
    return out


GROUP_FOR_SHEET = {"hparam-lr": "hparam"}   # sheet Summary chỉ có các nhóm chuẩn

COLUMNS = ["exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch",
           "epochs", "hidden", "dropout", "clip_norm", "precision", "init", "seed",
           "step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss",
           "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB", "diverged",
           "eval_acc", "eval_macro_f1", "figure_file", "notes"]

GROUP_ORDER = ["baseline", "loss", "optimizer", "hparam", "dropout", "clipping", "amp", "init",
               "final", "other"]


def _num(x, nd=6):
    if x is None:
        return None
    if isinstance(x, float) and (math.isnan(x) or math.isinf(x)):
        return None
    return round(x, nd) if isinstance(x, float) else x


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Một dòng của sheet Experiments, tên khoá trùng tên cột của mẫu.

    eval_scores: nội dung eval_result.json (chỉ cho baseline và cấu hình cuối), hoặc None.
    """
    c, s = result["cfg"], result["summary"]
    group = GROUP_FOR_SHEET.get(c.get("group"), c.get("group") or "other")
    extra = []
    if c.get("group") in GROUP_FOR_SHEET:
        extra.append(f"nhóm gốc: {c['group']}")
    if c.get("scheduler") not in (None, "none"):
        extra.append(f"scheduler={c['scheduler']}")
    if s.get("clip_frac"):
        extra.append(f"tỉ lệ bước bị clip={s['clip_frac']:.1%}")
    if s.get("diverged"):
        extra.append(f"diverged tại {s.get('diverged_at')}")
    if s.get("final_val_macro_f1") is not None and s.get("best_epoch") is not None:
        extra.append(f"val mF1 epoch cuối={s['final_val_macro_f1']:.4f}")
    all_notes = "; ".join([n for n in [notes] + extra if n])
    return {
        "exp_id": c["exp_id"],
        "group": group,
        "description": c.get("description", ""),
        "loss": c.get("loss"),
        "optimizer": c.get("optimizer"),
        "lr": c.get("lr"),
        "weight_decay": c.get("weight_decay"),
        "batch": c.get("batch"),
        "epochs": c.get("epochs"),
        "hidden": "-".join(str(h) for h in c.get("hidden", [])),
        "dropout": c.get("dropout"),
        "clip_norm": "none" if c.get("clip_norm") is None else c.get("clip_norm"),
        "precision": c.get("precision"),
        "init": c.get("init"),
        "seed": c.get("seed"),
        "step0_loss": _num(s.get("step0_loss")),
        "best_val_loss": _num(s.get("best_val_loss")),
        "best_epoch": s.get("best_epoch"),
        "final_train_loss": _num(s.get("final_train_loss")),
        "final_val_loss": _num(s.get("final_val_loss")),
        "val_acc": _num(s.get("val_acc")),
        "val_macro_f1": _num(s.get("val_macro_f1")),
        "time_per_epoch_s": _num(s.get("time_per_epoch_s"), 3),
        "peak_mem_MB": _num(s.get("peak_mem_MB"), 1),
        "diverged": "Có" if s.get("diverged") else "Không",
        "eval_acc": _num(eval_scores["accuracy"]) if eval_scores else None,
        "eval_macro_f1": _num(eval_scores["macro_f1"]) if eval_scores else None,
        "figure_file": f"{c['exp_id']}.png",
        "notes": all_notes,
    }


def sort_rows(rows: list[dict]) -> list[dict]:
    """Sắp theo thứ tự nhóm như sheet Summary, trong nhóm giữ thứ tự exp_id."""
    def key(r):
        g = r["group"]
        return (GROUP_ORDER.index(g) if g in GROUP_ORDER else len(GROUP_ORDER), r["exp_id"])
    return sorted(rows, key=key)


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền experiments.xlsx từ mẫu: giữ nguyên 4 sheet, tên cột và các cột công thức (AD..AG).

    rows          : list dict từ to_row (mỗi thí nghiệm một dòng)
    seed_ids      : exp_id các lần chạy baseline khác seed -> sheet Seeds cột A (tối đa 5)
    summary_notes : {group: nhận xét ngắn} -> sheet Summary cột H
    """
    from copy import copy
    from openpyxl import load_workbook

    wb = load_workbook(template_path)
    ws = wb["Experiments"]
    header = [cell.value for cell in ws[1]]
    col_of = {name: i + 1 for i, name in enumerate(header) if name}
    missing = [c for c in COLUMNS if c not in col_of]
    assert not missing, f"Mẫu thiếu cột: {missing}"

    formula_cols = [i + 1 for i, name in enumerate(header)
                    if isinstance(ws.cell(row=2, column=i + 1).value, str)
                    and str(ws.cell(row=2, column=i + 1).value).startswith("=")]
    last_template_row = ws.max_row

    # Xoá dữ liệu cũ (ô nhập), giữ nguyên định dạng và công thức
    for r in range(2, max(last_template_row, len(rows) + 1) + 1):
        for name in COLUMNS:
            ws.cell(row=r, column=col_of[name]).value = None

    for i, row in enumerate(rows):
        r = i + 2
        for name in COLUMNS:
            cell = ws.cell(row=r, column=col_of[name])
            cell.value = row.get(name)
            if r > last_template_row:                      # dòng vượt mẫu: chép định dạng dòng 2
                src = ws.cell(row=2, column=col_of[name])
                cell._style = copy(src._style)
        if r > last_template_row:                          # chép công thức cho dòng mới
            for col in formula_cols:
                f = ws.cell(row=2, column=col).value
                ws.cell(row=r, column=col).value = _shift_formula(f, 2, r)
                ws.cell(row=r, column=col)._style = copy(ws.cell(row=2, column=col)._style)

    if seed_ids is not None:
        ss = wb["Seeds"]
        for k in range(5):
            ss.cell(row=2 + k, column=1).value = seed_ids[k] if k < len(seed_ids) else None

    if summary_notes:
        sm = wb["Summary"]
        for r in range(2, sm.max_row + 1):
            g = sm.cell(row=r, column=1).value
            if g in summary_notes:
                sm.cell(row=r, column=8).value = summary_notes[g]

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)
    print(f"Đã ghi {len(rows)} dòng -> {out_path}")


def _shift_formula(formula: str, src_row: int, dst_row: int) -> str:
    """Dời tham chiếu ô tương đối (vd P2, T2) từ dòng src_row sang dst_row; giữ nguyên $C$8, $B$2:$B$61."""
    import re
    def repl(m):
        col, dollar_row, row = m.group(1), m.group(2), m.group(3)
        if dollar_row or int(row) != src_row:
            return m.group(0)
        return f"{col}{dst_row}"
    return re.sub(r"(?<![A-Za-z$!])([A-Z]{1,2})(\$?)(\d+)(?!\d)", repl, formula)